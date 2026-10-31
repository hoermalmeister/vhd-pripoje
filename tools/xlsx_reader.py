"""Nízkoúrovňové čtení listů z Excelu 'Přípoje mezi vlaky osobní dopravy'.

Každá buňka se převede na seznam segmentů (text, přeškrtnuto, červeně, font SenaKJR),
ze kterých se odvodí:
  effective  – platný text (bez přeškrtnutých částí, symboly KJŘ převedené na Unicode)
  original   – text před změnou (přeškrtnuté + nečervené části)
"""
import re, warnings
import openpyxl
from openpyxl.cell.rich_text import CellRichText, TextBlock

warnings.filterwarnings("ignore")

# Font SenaKJR (značky knižního jízdního řádu) -> Unicode
SENA_MAP = {"@": "①", "A": "②", "B": "③", "C": "④", "D": "⑤", "E": "⑥", "F": "⑦",
            "W": "⚒", "X": "†", "s": "–"}
RED = {"FFFF0000"}


def _is_red(color):
    try:
        return color is not None and isinstance(color.rgb, str) and color.rgb.upper() in RED
    except Exception:
        return False


def _seg(text, font_name, strike, red):
    if text is None:
        return None
    text = str(text)
    if font_name and "sena" in font_name.lower():
        text = "".join(SENA_MAP.get(ch, ch) for ch in text)
    return {"text": text, "strike": bool(strike), "red": bool(red)}


def cell_segments(cell):
    v = cell.value
    if v is None:
        return []
    f = cell.font
    segs = []
    if isinstance(v, CellRichText):
        for b in v:
            if isinstance(b, TextBlock):
                fn = b.font.rFont or f.name
                st = b.font.strike if b.font.strike is not None else f.strike
                segs.append(_seg(b.text, fn, st, _is_red(b.font.color)))
            else:
                segs.append(_seg(b, f.name, f.strike, _is_red(f.color)))
    else:
        if isinstance(v, float) and v.is_integer():
            v = int(v)
        if hasattr(v, "strftime"):
            v = v.strftime("%d.%m.%Y")
        segs.append(_seg(v, f.name, f.strike, _is_red(f.color)))
    return [s for s in segs if s and s["text"] != ""]


def norm_space(s):
    return re.sub(r"[ \t\r\n ]+", " ", s or "").strip()


class Cell:
    def __init__(self, segs):
        self.segs = segs
        self.effective = norm_space("".join(s["text"] for s in segs if not s["strike"]))
        self.original = norm_space("".join(s["text"] for s in segs if s["strike"] or not s["red"]))
        nonblank = [s for s in segs if s["text"].strip()]
        self.heuristic = None
        # '~~5~~ ~~-~~' : přeškrtnutí přes pomlčku není v PDF vidět -> pomlčka je nová hodnota
        if len(nonblank) >= 2 and all(s["strike"] for s in nonblank) and nonblank[-1]["text"].strip() in ("-", "–", "—") and nonblank[-1]["red"]:
            nonblank[-1]["strike"] = False
            self.heuristic = "dash_after_strike"
            self.effective = norm_space("".join(s["text"] for s in segs if not s["strike"]))
        self.fully_struck = bool(nonblank) and all(s["strike"] for s in nonblank)
        self.has_strike = any(s["strike"] and s["text"].strip() for s in segs)
        self.has_red = any(s["red"] and s["text"].strip() for s in segs)
        self.changed = self.has_strike or self.has_red

    @property
    def empty(self):
        return not any(s["text"].strip() for s in self.segs)

    def change_info(self):
        if not self.changed:
            return None
        d = {"before": self.original or None, "after": self.effective or None}
        if self.heuristic:
            d["heuristic"] = "Přeškrtnutí zasahuje i pomlčku, v PDF je pomlčka vidět jako platná – použita pomlčka."
        return d


def load_workbook(path):
    return openpyxl.load_workbook(path, rich_text=True)


def sheet_grid(ws, max_col=6):
    """Vrátí dict row -> [Cell x max_col]; svislé sloučení buněk se propíše dolů."""
    fill = {}
    for rng in ws.merged_cells.ranges:
        if rng.max_row > rng.min_row and rng.min_col <= max_col:
            top = ws.cell(rng.min_row, rng.min_col)
            for r in range(rng.min_row + 1, rng.max_row + 1):
                for c in range(rng.min_col, min(rng.max_col, max_col) + 1):
                    fill[(r, c)] = top if c == rng.min_col else None
    rows = {}
    for r in range(1, ws.max_row + 1):
        cells = []
        for c in range(1, max_col + 1):
            src = fill.get((r, c), ws.cell(r, c)) if (r, c) in fill else ws.cell(r, c)
            cell = Cell(cell_segments(src)) if src is not None else Cell([])
            cell.from_merge = (r, c) in fill
            cells.append(cell)
        rows[r] = cells
    return rows
